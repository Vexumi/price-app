# -*- coding: utf-8 -*-
"""
Помощник по Excel для ценового среза. НИКУДА НЕ ОБРАЩАЕТСЯ — только читает и
пишет таблицу. Фото распознаёт сам Claude Code; этот скрипт делает за него
надёжную часть: выдаёт список товаров по сети и проставляет цены в столбец-дату.

Команды:
  python excel_tools.py chains
      показать соответствие «папка -> сеть».

  python excel_tools.py candidates "Монетка" [--table таблица.xlsx]
      выдать товары этой сети: строка<TAB>продукт<TAB>производитель<TAB>масса.

  python excel_tools.py write 05.03.2026 matches.json [--table таблица.xlsx]
      создать (или дополнить) столбец-дату и проставить цены ПРЯМО в
      --table (пишет в этот же файл, копий не создаёт).
      matches.json: [{"row": 368, "price": 69.99, "discount": true}, ...]
      "discount": true — цена взята с акционного/скидочного ценника (было
      две цены, выбрали крупную); такая ячейка закрашивается жёлтым.
      Необязательно, по умолчанию false.
      Повторные вызовы за ту же дату дополняют тот же столбец в том же
      файле — можно обрабатывать сети по очереди.
"""
import argparse, json, datetime, copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill

DISCOUNT_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

FOLDER_TO_CHAIN = {
    "пятерочка":          "ПЯТЕРОЧКА",
    "магнит супермаркет": "МАГНИТ (S)",
    "магнит моя цена":    "МАГНИТ МОЯ ЦЕНА",
    "лента супермаркет":  "ЛЕНТА (S)",
    "лента гипермаркет":  "ЛЕНТА (G)",
    "монетка":            "МОНЕТКА",
    "мария-ра":           "МАРИЯ-РА",
    "ашан":               "АШАН",
    "окей":               "ОКЕЙ",
    "быстроном":          "БЫСТРОНОМ",
    "ярче":               "ЯРЧЕ",
    "чижик":              "ЧИЖИК",
}


def norm(v):
    return str(v).strip().upper() if v is not None else ""


def load_table(path):
    wb = load_workbook(path)
    ws = wb.active
    header_row = None
    for r in range(1, 15):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value).strip() == "Сеть":
                header_row = r
                break
        if header_row:
            break
    if not header_row:
        raise SystemExit("Не нашёл строку-заголовок со столбцом 'Сеть'.")
    cols, last_date = {}, None
    scan_max = min(ws.max_column, 400)        # реальных столбцов немного; справа бывает «мусор»
    for c in range(1, scan_max + 1):
        h = ws.cell(header_row, c).value
        hs = str(h).strip() if h is not None else ""
        if hs == "Сеть":             cols["chain"] = c
        elif hs == "Производитель":   cols["manuf"] = c
        elif hs == "Торговая марка":  cols["brand"] = c
        elif hs == "Продукт":         cols["product"] = c
        elif hs.startswith("Масса"):  cols["mass"] = c
        if isinstance(h, (datetime.datetime, datetime.date)):
            last_date = c
    cols["last_date"] = last_date
    for need in ("chain", "manuf", "product", "mass"):
        if need not in cols:
            raise SystemExit(f"Не нашёл столбец '{need}'.")
    return wb, ws, header_row, cols


def resolve_chain(name):
    return FOLDER_TO_CHAIN.get(name.strip().lower(), name.strip().upper())


def build_candidates(ws, header_row, cols, chain_canonical):
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        if norm(ws.cell(r, cols["chain"]).value) != chain_canonical:
            continue
        prod = ws.cell(r, cols["product"]).value
        if prod is None or str(prod).strip() == "":
            continue
        out.append((r, str(prod).strip(),
                    str(ws.cell(r, cols["manuf"]).value or "").strip(),
                    ws.cell(r, cols["mass"]).value))
    return out


def get_or_create_date_col(ws, header_row, cols, the_date):
    for c in range(1, min(ws.max_column, 400) + 1):
        h = ws.cell(header_row, c).value
        if isinstance(h, (datetime.datetime, datetime.date)):
            d = h.date() if isinstance(h, datetime.datetime) else h
            if d == the_date:
                return c
    new_col = (cols["last_date"] or cols["mass"]) + 1
    hdr = ws.cell(header_row, new_col,
                  datetime.datetime(the_date.year, the_date.month, the_date.day))
    if cols["last_date"]:
        src = ws.cell(header_row, cols["last_date"])
        hdr.font = copy.copy(src.font); hdr.fill = copy.copy(src.fill)
        hdr.border = copy.copy(src.border); hdr.alignment = copy.copy(src.alignment)
        hdr.number_format = src.number_format
        ws.column_dimensions[get_column_letter(new_col)].width = \
            ws.column_dimensions[get_column_letter(cols["last_date"])].width
    else:
        hdr.number_format = "DD.MM.YYYY"
    cols["last_date"] = new_col
    return new_col


def cmd_chains(_):
    print("папка  ->  значение столбца 'Сеть'")
    for k, v in FOLDER_TO_CHAIN.items():
        print(f"  {k:20} -> {v}")


def cmd_candidates(a):
    wb, ws, hr, cols = load_table(a.table)
    chain = resolve_chain(a.chain)
    cand = build_candidates(ws, hr, cols, chain)
    print(f"# сеть: {chain} | товаров: {len(cand)}")
    print("# row\tпродукт\tпроизводитель\tмасса")
    for row, prod, manuf, mass in cand:
        print(f"{row}\t{prod}\t{manuf}\t{mass}")


def cmd_write(a):
    the_date = datetime.datetime.strptime(a.date, "%d.%m.%Y").date()
    wb, ws, hr, cols = load_table(a.table)
    col = get_or_create_date_col(ws, hr, cols, the_date)
    with open(a.matches, encoding="utf-8-sig") as f:
        matches = json.load(f)
    n = 0
    for m in matches:
        cell = ws.cell(int(m["row"]), col, float(m["price"]))
        if m.get("discount"):
            cell.fill = DISCOUNT_FILL
        n += 1
    wb.save(a.table)
    print(f"Записано цен: {n} -> столбец {get_column_letter(col)} ({the_date:%d.%m.%Y}) в {a.table}")


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    s1 = sub.add_parser("chains"); s1.set_defaults(func=cmd_chains)
    s2 = sub.add_parser("candidates"); s2.add_argument("chain")
    s2.add_argument("--table", default="таблица.xlsx"); s2.set_defaults(func=cmd_candidates)
    s3 = sub.add_parser("write"); s3.add_argument("date"); s3.add_argument("matches")
    s3.add_argument("--table", default="таблица.xlsx")
    s3.set_defaults(func=cmd_write)
    a = p.parse_args(); a.func(a)


if __name__ == "__main__":
    main()
