# -*- coding: utf-8 -*-
"""Проверка матчинга товара на ручном эталоне."""
import json, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook
from price_app.parse import parse_photo
from price_app.match import find, decide, chain_of, mass_of

TABLE = "итоговый ценовой срез(АВГУСТ).xlsx"
# строка в таблице, установленная вручную по фото; None = товара в таблице нет
ET = {
 "Изображение PNG.png": 3724, "Изображение PNG 2.png": 181,
 "Изображение PNG 3.png": 635, "Изображение PNG 4.png": None,
 "Изображение PNG 5.png": 2494, "Изображение PNG 6.png": 1568,
 "Изображение PNG 7.png": "1632|1654", "Изображение PNG 8.png": None,
 "Изображение PNG 9.png": 254,
}

wb = load_workbook(TABLE); ws = wb.active
cands = {}
for r in range(6, ws.max_row + 1):
    ch = str(ws.cell(r, 1).value or "").strip().upper()
    prod = ws.cell(r, 6).value
    if not ch or not prod: continue
    cands.setdefault(ch, []).append(
        (r, str(prod), str(ws.cell(r, 2).value or ""), str(ws.cell(r, 3).value or ""),
         ws.cell(r, 11).value))

raw = json.load(open("out/ocr_tuned.json", encoding="utf-8"))
hit = miss = refused = 0
for name, want in ET.items():
    rec = raw[name]
    r = parse_photo(rec)
    full = " ".join(b["text"] for b in rec["blocks"])
    chain = chain_of(full) or "ЯРЧЕ"
    ranked = find(r["text"] or full, cands.get(chain, []))
    got, why = decide(ranked)
    gr = got[1] if got else None
    if want is None:
        ok = gr is None; verd = "верно отказался" if ok else f"ЛОЖНО взял {gr}"
    elif isinstance(want, str):
        ok = gr is not None and str(gr) in want.split("|")
        verd = "верно (спорный)" if ok else ("отказ" if gr is None else f"НЕВЕРНО {gr}")
    else:
        ok = gr == want
        verd = "верно" if ok else ("отказ" if gr is None else f"НЕВЕРНО {gr}")
    if ok: hit += 1
    elif gr is None: refused += 1
    else: miss += 1
    top = ", ".join(f"{s}:{rw}" for s, rw, _, _ in ranked[:3])
    print(f"{(name[14:-4] or 'PNG'):>5} | надо {str(want):>9} | {verd:<18} | {why:<28} | {top}")
print(f"\nверно: {hit}/{len(ET)} | НЕВЕРНО: {miss} | отказ: {refused}")
