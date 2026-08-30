# -*- coding: utf-8 -*-
"""Сводный прогон: цена + сеть + строка таблицы по всем фото."""
import json, sys, pathlib, collections
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook
from price_app.parse import parse_photo
from price_app.match import find, decide, chain_of

wb = load_workbook("итоговый ценовой срез(АВГУСТ).xlsx"); ws = wb.active
cands = collections.defaultdict(list)
for r in range(6, ws.max_row + 1):
    ch = str(ws.cell(r, 1).value or "").strip().upper()
    prod = ws.cell(r, 6).value
    if ch and prod:
        cands[ch].append((r, str(prod), str(ws.cell(r, 2).value or ""),
                          str(ws.cell(r, 3).value or ""), ws.cell(r, 11).value))

raw = json.load(open("out/ocr_328.json", encoding="utf-8"))
st = collections.Counter(); chains = collections.Counter()
for name, rec in sorted(raw.items()):
    r = parse_photo(rec)
    full = " ".join(b["text"] for b in rec["blocks"])
    chain = chain_of(full)
    chains[chain or "не определена"] += 1
    row = None; why = "-"
    if r["tag"]:
        pool = cands.get(chain, []) if chain else []
        if pool:
            got, why = decide(find(r["text"] or full, pool))
            row = got[1] if got else None
    if not r["tag"]:                     st["не ценник"] += 1
    elif r["price"] is not None and row: st["АВТОМАТ (цена+товар)"] += 1
    elif r["price"] is not None:         st["цена есть, товар на выбор"] += 1
    elif row:                            st["товар есть, цена на выбор"] += 1
    else:                                st["всё на проверку"] += 1
    p = f"{r['price']:7.2f}" if r["price"] else "    ---"
    print(f"{(name[14:-4] or 'PNG'):>5} | {str(chain or '?'):<11} | {p} | "
          f"{str(row or '-'):>5} | {why[:34]:<34} | {r['text'][:34]}")

n = sum(st.values())
print("\n" + "=" * 60)
for k, v in st.most_common(): print(f"  {v:3} ({100*v//n:2}%)  {k}")
print("\nсети:", dict(chains))
